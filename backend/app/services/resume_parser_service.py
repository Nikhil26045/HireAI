import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

import spacy
from sqlalchemy.orm import Session

from app.models.parsed_resume import ParsedResume
from app.models.resume import (
    Resume,
    ResumeProcessingStatus,
)


# ============================================================
# spaCy
# ============================================================

nlp = spacy.load("en_core_web_sm")


# ============================================================
# Technical Skills
# ============================================================

TECHNICAL_SKILLS = {
    # Programming languages
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "go",
    "rust",
    "php",
    "ruby",
    "kotlin",
    "swift",
    "r",
    "matlab",
    "verilog",
    "vhdl",
    "scala",

    # Web
    "html",
    "css",
    "react",
    "react.js",
    "next.js",
    "angular",
    "vue",
    "node.js",
    "express",
    "fastapi",
    "flask",
    "django",
    "bootstrap",
    "tailwind",
    "rest api",
    "rest",
    "graphql",
    "microservices",

    # Databases
    "sql",
    "postgresql",
    "mysql",
    "mongodb",
    "redis",
    "oracle",
    "sqlite",
    "mariadb",
    "database",

    # Python / Data
    "numpy",
    "pandas",
    "matplotlib",
    "seaborn",
    "scipy",
    "jupyter",

    # Machine Learning / AI
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "nlp",
    "natural language processing",
    "tensorflow",
    "pytorch",
    "scikit-learn",
    "sklearn",
    "spacy",
    "bert",
    "whisper",
    "gpt",
    "gpt-2",
    "llm",
    "llms",

    # Backend / ORM
    "sqlalchemy",
    "alembic",
    "pydantic",
    "uvicorn",

    # Cloud / DevOps
    "docker",
    "kubernetes",
    "git",
    "github",
    "gitlab",
    "bitbucket",
    "aws",
    "azure",
    "gcp",

    # Salesforce
    "salesforce",
    "apex",
    "soql",

    # Analytics / BI
    "tableau",
    "power bi",
    "powerbi",
    "etl",
    "data analytics",
    "data analysis",
    "business intelligence",

    # Engineering / Embedded
    "embedded systems",
    "stm32",
    "cortex-m4",
    "qemu",
    "gdb",
    "ti ccs",
    "ltspice",
    "octave",
    "xammp",
    "xampp",

    # Other tools
    "figma",
    "latex",
    "claude cli",
    "microsoft excel",
    "excel",
    "vba",
}


# ============================================================
# Skill aliases
# Canonical skill -> possible resume representations
# ============================================================

SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "JavaScript": ["javascript", "js"],
    "TypeScript": ["typescript", "ts"],
    "C": ["c"],
    "C++": ["c++", "cpp"],
    "C#": ["c#"],
    "Go": ["go"],
    "Rust": ["rust"],
    "PHP": ["php"],
    "Ruby": ["ruby"],
    "Kotlin": ["kotlin"],
    "Swift": ["swift"],
    "R": ["r"],
    "MATLAB": ["matlab"],
    "Verilog": ["verilog"],
    "VHDL": ["vhdl"],

    "HTML": ["html"],
    "CSS": ["css"],
    "React": ["react", "react.js"],
    "Next.js": ["next.js", "nextjs"],
    "Angular": ["angular"],
    "Vue": ["vue"],
    "Node.js": ["node.js", "nodejs"],
    "Express": ["express"],
    "FastAPI": ["fastapi"],
    "Flask": ["flask"],
    "Django": ["django"],
    "Bootstrap": ["bootstrap"],
    "Tailwind": ["tailwind"],

    "SQL": ["sql"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MySQL": ["mysql"],
    "MongoDB": ["mongodb"],
    "Redis": ["redis"],
    "Oracle": ["oracle"],
    "SQLite": ["sqlite"],

    "NumPy": ["numpy"],
    "Pandas": ["pandas"],
    "Matplotlib": ["matplotlib"],
    "Seaborn": ["seaborn"],
    "SciPy": ["scipy"],
    "Jupyter": ["jupyter"],

    "Machine Learning": ["machine learning"],
    "Deep Learning": ["deep learning"],
    "Artificial Intelligence": [
        "artificial intelligence",
        "artificial-intelligence",
    ],
    "NLP": [
        "nlp",
        "natural language processing",
    ],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"],
    "scikit-learn": [
        "scikit-learn",
        "sklearn",
    ],
    "spaCy": ["spacy"],
    "BERT": ["bert"],
    "Whisper": ["whisper"],
    "GPT-2": ["gpt-2", "gpt 2"],
    "LLM": ["llm", "llms"],

    "SQLAlchemy": ["sqlalchemy"],
    "Alembic": ["alembic"],
    "Pydantic": ["pydantic"],
    "Uvicorn": ["uvicorn"],

    "Docker": ["docker"],
    "Kubernetes": ["kubernetes"],
    "Git": ["git"],
    "GitHub": ["github"],
    "GitLab": ["gitlab"],
    "Bitbucket": ["bitbucket"],
    "AWS": ["aws"],
    "Azure": ["azure"],
    "GCP": ["gcp"],

    "Salesforce": ["salesforce"],
    "Apex": ["apex"],
    "SOQL": ["soql"],

    "REST API": ["rest api"],
    "REST": ["rest"],
    "GraphQL": ["graphql"],
    "Microservices": ["microservices"],

    "Tableau": ["tableau"],
    "Power BI": ["power bi", "powerbi"],
    "ETL": ["etl"],
    "Data Analytics": ["data analytics"],
    "Data Analysis": ["data analysis"],
    "Business Intelligence": ["business intelligence"],

    "Embedded Systems": ["embedded systems"],
    "STM32": ["stm32"],
    "Cortex-M4": ["cortex-m4", "cortex-m4f"],
    "QEMU": ["qemu"],
    "GDB": ["gdb"],
    "TI CCS": ["ti ccs"],
    "LTSpice": ["ltspice"],
    "Octave": ["octave"],
    "XAMPP": ["xampp", "xammp"],

    "Figma": ["figma"],
    "LaTeX": ["latex"],
    "Claude CLI": ["claude cli"],
    "Excel": ["excel"],
    "VBA": ["vba"],
}


# ============================================================
# Resume Section Aliases
# ============================================================

SECTION_ALIASES = {
    "summary": {
        "summary",
        "professional summary",
        "profile",
        "objective",
        "career objective",
        "about me",
        "about",
    },

    "skills": {
        "skills",
        "technical skills",
        "technical expertise",
        "technical competencies",
        "core skills",
        "technologies",
        "technical proficiency",
    },

    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "work history",
        "internships",
        "internship",
    },

    "education": {
        "education",
        "academic background",
        "academic qualifications",
        "educational qualifications",
        "education qualifications",
    },

    "projects": {
        "projects",
        "key projects",
        "personal projects",
        "academic projects",
        "project experience",
        "key project",
    },

    "certifications": {
        "certifications",
        "certificates",
        "licenses & certifications",
        "certification",
    },

    "achievements": {
        "achievements",
        "scholastic achievements",
        "academic achievements",
        "scholastic",
        "awards",
        "honors",
        "honours",
    },

    "positions": {
        "positions of responsibility",
        "positions of responsibilities",
        "position of responsibility",
        "leadership",
        "leadership experience",
        "responsibilities",
    },

    "extracurricular": {
        "extra curricular activities",
        "extracurricular activities",
        "extracurricular",
        "extra curricular",
        "activities",
    },

    "coursework": {
        "relevant courses",
        "relevant coursework",
        "coursework",
        "academic courses",
    },
}


# ============================================================
# Heading normalization
# Handles PDF extraction such as:
#
# TECHNICALSKILLS
# KEYPROJECTS
# POSITIONSOFRESPONSIBILITIES
# ============================================================

HEADING_NORMALIZATION = {
    "summary": "summary",
    "professionalsummary": "professional summary",
    "profile": "profile",
    "objective": "objective",
    "careerobjective": "career objective",
    "aboutme": "about me",
    "about": "about",

    "skills": "skills",
    "technicalskills": "technical skills",
    "technicalexpertise": "technical expertise",
    "technicalcompetencies": "technical competencies",
    "coreskills": "core skills",
    "technologies": "technologies",
    "technicalproficiency": "technical proficiency",

    "experience": "experience",
    "workexperience": "work experience",
    "professionalexperience": "professional experience",
    "employmenthistory": "employment history",
    "workhistory": "work history",
    "internships": "internships",
    "internship": "internship",

    "education": "education",
    "academicbackground": "academic background",
    "academicqualifications": "academic qualifications",
    "educationalqualifications": "educational qualifications",
    "educationqualifications": "education qualifications",

    "projects": "projects",
    "keyprojects": "key projects",
    "personalprojects": "personal projects",
    "academicprojects": "academic projects",
    "projectexperience": "project experience",
    "keyproject": "key project",

    "certifications": "certifications",
    "certificates": "certificates",
    "licensescertifications": "licenses & certifications",
    "certification": "certification",

    "achievements": "achievements",
    "scholasticachievements": "scholastic achievements",
    "academicachievements": "academic achievements",
    "scholastic": "scholastic",
    "awards": "awards",
    "honors": "honors",
    "honours": "honours",

    "positionsofresponsibility": "positions of responsibility",
    "positionsofresponsibilities": "positions of responsibilities",
    "positionofresponsibility": "position of responsibility",
    "leadership": "leadership",
    "leadershipexperience": "leadership experience",
    "responsibilities": "responsibilities",

    "extracurricularactivities": "extracurricular activities",
    "extracurricular": "extracurricular",
    "extracurricularactivities": "extracurricular activities",
    "extracurricular": "extracurricular",

    "relevantcourses": "relevant courses",
    "relevantcoursework": "relevant coursework",
    "coursework": "coursework",
    "academiccourses": "academic courses",
}


# ============================================================
# Common patterns
# ============================================================

DATE_PATTERN = (
    r"(?:"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)"
    r"[a-z]*[’']?\d{2,4}"
    r"|"
    r"(?:19|20)\d{2}"
    r")"
)

DATE_RANGE_PATTERN = (
    rf"({DATE_PATTERN})"
    r"\s*(?:[-–—]|to)\s*"
    rf"({DATE_PATTERN}|present|current)"
)


# ============================================================
# Text normalization
# ============================================================

def normalize_text(text: str) -> str:
    """
    Normalize PDF/DOCX extracted text while preserving lines.

    Important:
    We intentionally do NOT remove newlines because resume
    section detection depends heavily on line boundaries.
    """

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Replace common PDF artifacts
    text = text.replace("\u00a0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\ufeff", "")

    lines = []

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        # Collapse repeated whitespace
        line = re.sub(r"[ \t]+", " ", line)

        # Normalize bullets
        line = re.sub(r"^[•▪●◦‣⁃]+\s*", "• ", line)

        lines.append(line)

    return "\n".join(lines)


# ============================================================
# PDF Extraction
# ============================================================

def extract_text_from_pdf(
    file_path: Path,
) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages).strip()


# ============================================================
# DOCX Extraction
# ============================================================

def extract_text_from_docx(
    file_path: Path,
) -> str:
    from docx import Document

    document = Document(str(file_path))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also capture text inside tables.
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                value = cell.text.strip()

                if value:
                    cells.append(value)

            if cells:
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs).strip()


# ============================================================
# Resume File Text Extraction
# ============================================================

def extract_resume_text(
    file_path: Path,
    file_type: str,
) -> str:

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    raise ValueError(
        "Unsupported resume file format. "
        "Only PDF and DOCX are currently supported."
    )


# ============================================================
# Utility Functions
# ============================================================

def clean_text(value: str | None) -> str | None:
    if value is None:
        return None

    value = value.replace("\u00a0", " ")
    value = re.sub(r"\s+", " ", value)
    value = value.strip()

    return value or None


def clean_bullet(value: str) -> str:
    value = re.sub(
        r"^[•▪●◦‣⁃\-–—*]+\s*",
        "",
        value,
    )

    return clean_text(value) or ""


def normalize_heading(line: str) -> str:
    normalized = line.lower().strip()

    normalized = normalized.rstrip(":")
    normalized = normalized.replace("&", "and")

    # Remove bullet characters
    normalized = re.sub(
        r"^[•▪●◦‣⁃\-–—*]+\s*",
        "",
        normalized,
    )

    # Remove punctuation and spaces for compact matching
    compact = re.sub(
        r"[^a-z0-9]+",
        "",
        normalized,
    )

    if compact in HEADING_NORMALIZATION:
        return HEADING_NORMALIZATION[compact]

    return normalized


def compact_text(value: str) -> str:
    return re.sub(
        r"[^a-z0-9]+",
        "",
        value.lower(),
    )


def looks_like_bullet(line: str) -> bool:
    return bool(
        re.match(
            r"^\s*[•▪●◦‣⁃\-–—*]",
            line,
        )
    )


def looks_like_heading(line: str) -> bool:
    """
    Detect common resume headings.

    Works for:
    TECHNICAL SKILLS
    TECHNICALSKILLS
    Key Projects
    KEYPROJECTS
    """

    cleaned = line.strip()

    if not cleaned:
        return False

    normalized = normalize_heading(cleaned)

    for aliases in SECTION_ALIASES.values():
        for alias in aliases:
            if normalized == alias:
                return True

    compact = compact_text(cleaned)

    if compact in HEADING_NORMALIZATION:
        return True

    # Strong visual-heading heuristic
    alpha_chars = [
        char for char in cleaned
        if char.isalpha()
    ]

    if len(alpha_chars) >= 5:
        uppercase_ratio = sum(
            char.isupper()
            for char in alpha_chars
        ) / len(alpha_chars)

        if uppercase_ratio >= 0.85:
            if len(cleaned.split()) <= 8:
                return True

    return False


# ============================================================
# Email Extraction
# ============================================================

def extract_email(
    text: str,
) -> str | None:

    pattern = (
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
    )

    match = re.search(
        pattern,
        text,
    )

    if match:
        return match.group(0).strip()

    return None


# ============================================================
# Phone Extraction
# ============================================================

def extract_phone(
    text: str,
) -> str | None:

    patterns = [
        r"\+\d{1,3}[\s\-()]?\d{3}[\s\-()]?\d{3}[\s\-()]?\d{4}",
        r"\+\d{1,3}[\s\-()]?\d{8,12}",
        r"\b\d{10}\b",
        r"\b\d{3}[\s\-]\d{3}[\s\-]\d{4}\b",
        r"\b\d{5}[\s\-]\d{5}\b",
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
        )

        if match:
            return clean_text(match.group(0))

    return None


# ============================================================
# Name Extraction
# ============================================================

def extract_name(
    text: str,
) -> str | None:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return None

    # --------------------------------------------------------
    # 1. Strongest case:
    # Standalone uppercase name.
    # --------------------------------------------------------

    for line in lines[:8]:

        candidate = re.sub(
            r"[^A-Za-z\s'.-]",
            "",
            line,
        ).strip()

        words = candidate.split()

        if not (2 <= len(words) <= 4):
            continue

        if candidate.upper() != candidate:
            continue

        lowered = candidate.lower()

        blocked = {
            "department",
            "engineering",
            "undergraduate",
            "student",
            "resume",
            "curriculum vitae",
            "education",
            "experience",
            "skills",
        }

        if any(
            item in lowered
            for item in blocked
        ):
            continue

        return candidate.title()

    # --------------------------------------------------------
    # 2. Name after "Undergraduate", "Graduate", etc.
    # --------------------------------------------------------

    patterns = [
        r"(?:fourth|third|second|first)\s+year\s+"
        r"(?:undergraduate|student)\s+"
        r"([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+){1,3})",

        r"(?:undergraduate|graduate|student)\s+"
        r"([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+){1,3})",
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE,
        )

        if match:
            candidate = clean_text(match.group(1))

            if candidate:
                return candidate

    # --------------------------------------------------------
    # 3. Look for a 2–4 word name near the top.
    # --------------------------------------------------------

    blocked_terms = {
        "fourth",
        "third",
        "second",
        "first",
        "year",
        "undergraduate",
        "graduate",
        "student",
        "department",
        "electrical",
        "engineering",
        "institute",
        "university",
        "college",
        "email",
        "phone",
        "linkedin",
        "github",
    }

    for line in lines[:10]:

        if any(
            char.isdigit()
            for char in line
        ):
            continue

        if "@" in line:
            continue

        candidate = re.sub(
            r"[^A-Za-z\s'.-]",
            "",
            line,
        ).strip()

        words = candidate.split()

        if not (2 <= len(words) <= 4):
            continue

        if any(
            word.lower() in blocked_terms
            for word in words
        ):
            continue

        return candidate.title()

    return None


# ============================================================
# Summary Extraction
# ============================================================

def extract_summary(
    text: str,
) -> str | None:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    summary_aliases = SECTION_ALIASES["summary"]

    for index, line in enumerate(lines):

        normalized = normalize_heading(line)

        if normalized not in summary_aliases:
            continue

        summary_lines = []

        for next_line in lines[index + 1:]:
            if looks_like_heading(next_line):
                break

            summary_lines.append(
                clean_bullet(next_line)
            )

        if summary_lines:
            return clean_text(
                " ".join(summary_lines)
            )

    return None


# ============================================================
# Experience Years Extraction
# ============================================================

def extract_experience_years(
    text: str,
) -> int | None:

    patterns = [
        r"(\d+)\+?\s+years?\s+(?:of\s+)?experience",
        r"experience\s*[:\-]?\s*(\d+)\+?\s+years?",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE,
        )

        if match:
            return int(match.group(1))

    return None


# ============================================================
# Date Extraction
# ============================================================

def extract_date_range(
    text: str,
) -> dict | None:

    match = re.search(
        DATE_RANGE_PATTERN,
        text,
        flags=re.IGNORECASE,
    )

    if not match:
        return None

    return {
        "start_date": clean_text(match.group(1)),
        "end_date": clean_text(match.group(2)),
        "raw": clean_text(match.group(0)),
    }


def remove_date_range(text: str) -> str:
    return clean_text(
        re.sub(
            DATE_RANGE_PATTERN,
            "",
            text,
            flags=re.IGNORECASE,
        )
    ) or ""


# ============================================================
# Section Detection
# ============================================================

def detect_section_heading(
    line: str,
) -> str | None:

    normalized = normalize_heading(line)

    for section, aliases in SECTION_ALIASES.items():

        if normalized in aliases:
            return section

    compact = compact_text(normalized)

    for section, aliases in SECTION_ALIASES.items():

        for alias in aliases:

            if compact == compact_text(alias):
                return section

    return None


# ============================================================
# Section Extraction
# ============================================================

def extract_sections(
    text: str,
) -> dict[str, list[str]]:

    sections: dict[str, list[str]] = {}

    current_section = "general"

    sections[current_section] = []

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        section = detect_section_heading(line)

        if section:

            current_section = section

            if current_section not in sections:
                sections[current_section] = []

            continue

        sections.setdefault(
            current_section,
            [],
        ).append(line)

    return {
        key: value
        for key, value in sections.items()
        if value
    }


# ============================================================
# Education Extraction
# ============================================================

def extract_education(
    text: str,
    sections: dict[str, list[str]],
) -> list[dict]:

    education = []

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    # --------------------------------------------------------
    # Detect table header.
    # --------------------------------------------------------

    table_started = False

    for line in lines:

        compact = compact_text(line)

        if (
            "examination" in compact
            and "degreecertificate" in compact
            and "institute" in compact
        ):
            table_started = True
            continue

        if not table_started:
            continue

        # Stop at an obvious resume section
        if detect_section_heading(line):
            break

        # ----------------------------------------------------
        # Known academic row patterns
        # ----------------------------------------------------

        row_match = re.match(
            r"^(Graduation|Intermediate/?\+?2|Matriculation)"
            r"\s+(.+?)\s+"
            r"(.+?)\s+"
            r"((?:19|20)\d{2}(?:\s*[–-]\s*(?:Present|"
            r"(?:19|20)\d{2}))?)"
            r"\s+"
            r"(\d+(?:\.\d+)?(?:/\d+|%))$",
            line,
            flags=re.IGNORECASE,
        )

        if row_match:

            education.append(
                {
                    "level": clean_text(
                        row_match.group(1)
                    ),
                    "degree": clean_text(
                        row_match.group(2)
                    ),
                    "institution": clean_text(
                        row_match.group(3)
                    ),
                    "year": clean_text(
                        row_match.group(4)
                    ),
                    "score": clean_text(
                        row_match.group(5)
                    ),
                }
            )

            continue

        # ----------------------------------------------------
        # Fallback for compressed rows.
        # ----------------------------------------------------

        if line.startswith(
            (
                "Graduation ",
                "Intermediate",
                "Matriculation ",
            )
        ):

            score_match = re.search(
                r"(\d+(?:\.\d+)?(?:/\d+|%))$",
                line,
            )

            if not score_match:
                continue

            score = score_match.group(1)

            before_score = line[
                :score_match.start()
            ].strip()

            year_match = re.search(
                r"((?:19|20)\d{2}"
                r"(?:\s*[–-]\s*(?:Present|"
                r"(?:19|20)\d{2}))?)",
                before_score,
                flags=re.IGNORECASE,
            )

            if not year_match:
                continue

            year = year_match.group(1)

            prefix = before_score[
                :year_match.start()
            ].strip()

            parts = prefix.split()

            if len(parts) < 3:
                continue

            level = parts[0]

            if level.lower().startswith("intermediate"):
                level = "Intermediate/+2"

            degree_candidates = {
                "b.tech",
                "btech",
                "cbse(xii)",
                "cbse(x)",
            }

            degree_index = None

            for index, part in enumerate(parts[1:], start=1):

                if part.lower() in degree_candidates:
                    degree_index = index
                    break

            if degree_index is None:
                continue

            degree = parts[degree_index]
            institution = " ".join(
                parts[degree_index + 1:]
            )

            education.append(
                {
                    "level": clean_text(level),
                    "degree": clean_text(degree),
                    "institution": clean_text(institution),
                    "year": clean_text(year),
                    "score": clean_text(score),
                }
            )

    # --------------------------------------------------------
    # If no education table was found, use education section.
    # --------------------------------------------------------

    if not education:

        education_lines = sections.get(
            "education",
            [],
        )

        for line in education_lines:

            education.append(
                {
                    "raw": clean_bullet(line),
                }
            )

    return education


# ============================================================
# Skills Extraction
# ============================================================

def extract_skills(
    text: str,
) -> list[str]:

    normalized_text = text.lower()

    found_skills = set()

    # --------------------------------------------------------
    # Search canonical aliases.
    # --------------------------------------------------------

    for canonical, aliases in SKILL_ALIASES.items():

        for alias in aliases:

            escaped = re.escape(
                alias.lower()
            )

            # Word-boundary matching.
            pattern = (
                rf"(?<!\w)"
                rf"{escaped}"
                rf"(?!\w)"
            )

            if re.search(
                pattern,
                normalized_text,
            ):
                found_skills.add(canonical)
                break

    # --------------------------------------------------------
    # Special case: C/C++
    # --------------------------------------------------------

    if re.search(
        r"\bc\s*/\s*c\+\+",
        normalized_text,
    ):
        found_skills.add("C")
        found_skills.add("C++")

    # --------------------------------------------------------
    # Special case: Git/GitHub
    # --------------------------------------------------------

    if re.search(
        r"\bgit\s*/\s*github\b",
        normalized_text,
    ):
        found_skills.add("Git")
        found_skills.add("GitHub")

    return sorted(
        found_skills,
        key=lambda value: value.lower(),
    )


# ============================================================
# Skills by Category
# ============================================================

def extract_categorized_skills(
    text: str,
) -> dict:

    result = {
        "languages": [],
        "libraries_frameworks": [],
        "databases": [],
        "tools": [],
        "other": [],
    }

    language_skills = {
        "Python",
        "Java",
        "JavaScript",
        "TypeScript",
        "C",
        "C++",
        "C#",
        "Go",
        "Rust",
        "PHP",
        "Ruby",
        "Kotlin",
        "Swift",
        "R",
        "MATLAB",
        "Verilog",
        "VHDL",
    }

    database_skills = {
        "SQL",
        "PostgreSQL",
        "MySQL",
        "MongoDB",
        "Redis",
        "Oracle",
        "SQLite",
    }

    library_skills = {
        "React",
        "Next.js",
        "Angular",
        "Vue",
        "Node.js",
        "Express",
        "FastAPI",
        "Flask",
        "Django",
        "NumPy",
        "Pandas",
        "Matplotlib",
        "Seaborn",
        "SciPy",
        "PyTorch",
        "TensorFlow",
        "scikit-learn",
        "spaCy",
        "BERT",
        "Whisper",
        "SQLAlchemy",
        "Alembic",
        "Pydantic",
    }

    tool_skills = {
        "Git",
        "GitHub",
        "GitLab",
        "Docker",
        "Kubernetes",
        "AWS",
        "Azure",
        "GCP",
        "Tableau",
        "Power BI",
        "Jupyter",
        "Figma",
        "QEMU",
        "GDB",
        "TI CCS",
        "LTSpice",
        "Octave",
        "XAMPP",
        "Claude CLI",
        "LaTeX",
        "Excel",
        "VBA",
    }

    all_skills = extract_skills(text)

    for skill in all_skills:

        if skill in language_skills:
            result["languages"].append(skill)

        elif skill in database_skills:
            result["databases"].append(skill)

        elif skill in library_skills:
            result["libraries_frameworks"].append(skill)

        elif skill in tool_skills:
            result["tools"].append(skill)

        else:
            result["other"].append(skill)

    return result


# ============================================================
# Generic Bullet Extraction
# ============================================================

def extract_bullets(
    lines: list[str],
) -> list[str]:

    bullets = []

    for line in lines:

        cleaned = clean_bullet(line)

        if not cleaned:
            continue

        bullets.append(cleaned)

    return bullets


# ============================================================
# Experience Extraction
# ============================================================

def is_probable_experience_header(
    line: str,
) -> bool:

    if looks_like_bullet(line):
        return False

    if detect_section_heading(line):
        return False

    # Date range is a strong signal.
    if re.search(
        DATE_RANGE_PATTERN,
        line,
        flags=re.IGNORECASE,
    ):
        return True

    return False


def extract_experience(
    sections: dict[str, list[str]],
) -> list[dict]:

    lines = sections.get(
        "experience",
        [],
    )

    experiences = []

    current = None
    pending_role = None

    for line in lines:

        cleaned = clean_bullet(line)

        if not cleaned:
            continue

        # ----------------------------------------------------
        # New company/experience header.
        # ----------------------------------------------------

        if is_probable_experience_header(line):

            date_info = extract_date_range(
                line
            )

            header_without_date = (
                remove_date_range(line)
            )

            header_without_date = clean_bullet(
                header_without_date
            )

            # Company may contain location.
            company = header_without_date
            location = None

            # Common pattern:
            # Company, City
            if "," in header_without_date:

                parts = [
                    part.strip()
                    for part in header_without_date.split(
                        ",",
                        maxsplit=1,
                    )
                ]

                if len(parts) == 2:
                    company = parts[0]
                    location = parts[1]

            current = {
                "company": clean_text(company),
                "location": clean_text(location),
                "role": None,
                "start_date": (
                    date_info["start_date"]
                    if date_info
                    else None
                ),
                "end_date": (
                    date_info["end_date"]
                    if date_info
                    else None
                ),
                "description": [],
            }

            experiences.append(current)
            pending_role = current

            continue

        # ----------------------------------------------------
        # If we don't have an experience yet, create one.
        # ----------------------------------------------------

        if current is None:

            current = {
                "company": None,
                "location": None,
                "role": None,
                "start_date": None,
                "end_date": None,
                "description": [],
            }

            experiences.append(current)

        # ----------------------------------------------------
        # Role line.
        # ----------------------------------------------------

        if (
            current["role"] is None
            and not looks_like_bullet(line)
        ):

            # Don't classify obvious dates as role.
            if not re.fullmatch(
                r"[\d\s/–—'-]+",
                cleaned,
            ):

                current["role"] = cleaned
                continue

        # ----------------------------------------------------
        # Description.
        # ----------------------------------------------------

        current["description"].append(
            cleaned
        )

    # --------------------------------------------------------
    # Remove empty placeholder records.
    # --------------------------------------------------------

    cleaned_experiences = []

    for experience in experiences:

        if not any(
            [
                experience.get("company"),
                experience.get("role"),
                experience.get("description"),
            ]
        ):
            continue

        cleaned_experiences.append(
            experience
        )

    return cleaned_experiences


# ============================================================
# Project Extraction
# ============================================================

def is_probable_project_header(
    line: str,
) -> bool:

    if looks_like_bullet(line):
        return False

    if detect_section_heading(line):
        return False

    if re.search(
        DATE_RANGE_PATTERN,
        line,
        flags=re.IGNORECASE,
    ):
        return True

    # Project titles often don't have dates.
    # Detect them when followed by a project-type/supervisor
    # line is handled by the parser state.
    return False


def extract_projects(
    sections: dict[str, list[str]],
) -> list[dict]:

    lines = sections.get(
        "projects",
        [],
    )

    projects = []

    current = None

    for index, line in enumerate(lines):

        cleaned = clean_bullet(line)

        if not cleaned:
            continue

        if looks_like_bullet(line):

            if current is None:
                current = {
                    "title": None,
                    "type": None,
                    "supervisor": None,
                    "start_date": None,
                    "end_date": None,
                    "description": [],
                }

                projects.append(current)

            current["description"].append(
                cleaned
            )

            continue

        # ----------------------------------------------------
        # Detect date range.
        # ----------------------------------------------------

        date_info = extract_date_range(
            line
        )

        # ----------------------------------------------------
        # If current project is missing title,
        # current non-bullet line is title.
        # ----------------------------------------------------

        if current is None:

            current = {
                "title": None,
                "type": None,
                "supervisor": None,
                "start_date": None,
                "end_date": None,
                "description": [],
            }

            projects.append(current)

            if date_info:
                current["title"] = (
                    remove_date_range(line)
                )

                current["start_date"] = (
                    date_info["start_date"]
                )

                current["end_date"] = (
                    date_info["end_date"]
                )
            else:
                current["title"] = cleaned

            continue

        # ----------------------------------------------------
        # Current project already has title.
        # ----------------------------------------------------

        if (
            current["title"]
            and current["type"] is None
            and not current["description"]
        ):

            lowered = cleaned.lower()

            if (
                "project" in lowered
                or "independent" in lowered
                or "b.tech" in lowered
                or "academic" in lowered
                or lowered.startswith("prof.")
            ):

                if lowered.startswith("prof."):
                    current["supervisor"] = cleaned
                else:
                    current["type"] = cleaned

                continue

        # ----------------------------------------------------
        # Supervisor line.
        # ----------------------------------------------------

        if (
            current["supervisor"] is None
            and (
                cleaned.lower().startswith("prof.")
                or "prof." in cleaned.lower()
            )
            and not current["description"]
        ):

            current["supervisor"] = cleaned
            continue

        # ----------------------------------------------------
        # If a new non-bullet line comes after a project
        # has content, it is likely the next project.
        # ----------------------------------------------------

        if current["description"]:

            current = {
                "title": None,
                "type": None,
                "supervisor": None,
                "start_date": None,
                "end_date": None,
                "description": [],
            }

            projects.append(current)

            if date_info:

                current["title"] = (
                    remove_date_range(line)
                )

                current["start_date"] = (
                    date_info["start_date"]
                )

                current["end_date"] = (
                    date_info["end_date"]
                )

            else:
                current["title"] = cleaned

            continue

        # ----------------------------------------------------
        # Additional metadata.
        # ----------------------------------------------------

        if current["type"] is None:
            current["type"] = cleaned

    # --------------------------------------------------------
    # Clean projects.
    # --------------------------------------------------------

    cleaned_projects = []

    for project in projects:

        if not project.get("title"):
            continue

        cleaned_projects.append(project)

    return cleaned_projects


# ============================================================
# Certification Extraction
# ============================================================

def extract_certifications(
    sections: dict[str, list[str]],
) -> list[str]:

    lines = sections.get(
        "certifications",
        [],
    )

    return extract_bullets(lines)


# ============================================================
# Achievement Extraction
# ============================================================

def extract_achievements(
    sections: dict[str, list[str]],
) -> list[str]:

    lines = sections.get(
        "achievements",
        [],
    )

    return extract_bullets(lines)


# ============================================================
# Positions of Responsibility
# ============================================================

def extract_positions(
    sections: dict[str, list[str]],
) -> list[dict]:

    lines = sections.get(
        "positions",
        [],
    )

    positions = []

    current = None

    for line in lines:

        cleaned = clean_bullet(line)

        if not cleaned:
            continue

        date_info = extract_date_range(
            line
        )

        # ----------------------------------------------------
        # New position header.
        # ----------------------------------------------------

        if not looks_like_bullet(line):

            header = (
                remove_date_range(line)
                if date_info
                else cleaned
            )

            current = {
                "position": clean_text(header),
                "start_date": (
                    date_info["start_date"]
                    if date_info
                    else None
                ),
                "end_date": (
                    date_info["end_date"]
                    if date_info
                    else None
                ),
                "description": [],
            }

            positions.append(current)

            continue

        if current is None:

            current = {
                "position": None,
                "start_date": None,
                "end_date": None,
                "description": [],
            }

            positions.append(current)

        current["description"].append(
            cleaned
        )

    return [
        position
        for position in positions
        if position.get("position")
        or position.get("description")
    ]


# ============================================================
# Extracurricular Activities
# ============================================================

def extract_extracurricular(
    sections: dict[str, list[str]],
) -> list[str]:

    lines = sections.get(
        "extracurricular",
        [],
    )

    return extract_bullets(lines)


# ============================================================
# Relevant Coursework
# ============================================================

def extract_coursework(
    sections: dict[str, list[str]],
) -> list[str]:

    lines = sections.get(
        "coursework",
        [],
    )

    coursework = []

    for line in lines:

        cleaned = clean_bullet(line)

        if not cleaned:
            continue

        # Remove footnote symbols.
        cleaned = re.sub(
            r"[†‡]",
            "",
            cleaned,
        )

        # PDF extraction frequently combines two or more
        # courses on one line. We preserve the complete line
        # rather than incorrectly splitting course names.
        coursework.append(
            clean_text(cleaned)
        )

    return [
        course
        for course in coursework
        if course
    ]


# ============================================================
# Header Information
# ============================================================

def extract_header(
    text: str,
) -> dict:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    header_lines = lines[:15]

    institution = None
    department = None
    current_position = None

    for line in header_lines:

        lowered = line.lower()

        if "department" in lowered:

            department = clean_text(
                line
            )

        if (
            "institute" in lowered
            or "university" in lowered
            or "college" in lowered
        ):

            if not institution:
                institution = clean_text(
                    line
                )

        if (
            "undergraduate" in lowered
            or "graduate" in lowered
            or "student" in lowered
        ):

            current_position = clean_text(
                line
            )

    return {
        "institution": institution,
        "department": department,
        "current_position": current_position,
    }


# ============================================================
# General Academic Details
# ============================================================

def extract_academic_profile(
    text: str,
) -> dict:

    result = {
        "degree": None,
        "institution": None,
        "department": None,
        "current_status": None,
        "cpi": None,
    }

    header = extract_header(text)

    result["institution"] = (
        header["institution"]
    )

    result["department"] = (
        header["department"]
    )

    result["current_status"] = (
        header["current_position"]
    )

    degree_match = re.search(
        r"\bB\.?\s*Tech\.?\b",
        text,
        flags=re.IGNORECASE,
    )

    if degree_match:
        result["degree"] = "B.Tech"

    cpi_match = re.search(
        r"(\d+(?:\.\d+)?)\s*/\s*10",
        text,
    )

    if cpi_match:
        result["cpi"] = cpi_match.group(1)

    return result


# ============================================================
# Better Named Entity Extraction
# ============================================================

def extract_named_entities(
    text: str,
) -> list[dict]:

    doc = nlp(text)

    entities = []

    allowed_labels = {
        "PERSON",
        "ORG",
        "GPE",
        "FAC",
    }

    seen = set()

    for entity in doc.ents:

        label = entity.label_
        value = clean_text(entity.text)

        if not value:
            continue

        if label not in allowed_labels:
            continue

        # Ignore very noisy fragments.
        if len(value) < 3:
            continue

        if value.startswith(
            (
                "•",
                "-",
                "–",
                "—",
            )
        ):
            continue

        key = (
            value.lower(),
            label,
        )

        if key in seen:
            continue

        seen.add(key)

        entities.append(
            {
                "text": value,
                "label": label,
            }
        )

    return entities


# ============================================================
# Extract Important Organizations
# ============================================================

def extract_organizations(
    text: str,
) -> list[str]:

    doc = nlp(text)

    organizations = []
    seen = set()

    for entity in doc.ents:

        if entity.label_ not in {
            "ORG",
            "FAC",
        }:
            continue

        value = clean_text(entity.text)

        if not value:
            continue

        key = value.lower()

        if key in seen:
            continue

        seen.add(key)
        organizations.append(value)

    return organizations


# ============================================================
# Extract People
# ============================================================

def extract_people(
    text: str,
) -> list[str]:

    doc = nlp(text)

    people = []
    seen = set()

    for entity in doc.ents:

        if entity.label_ != "PERSON":
            continue

        value = clean_text(entity.text)

        if not value:
            continue

        # Ignore very noisy fragments.
        if len(value.split()) > 6:
            continue

        key = value.lower()

        if key in seen:
            continue

        seen.add(key)
        people.append(value)

    return people


# ============================================================
# Detect Resume Sections
# ============================================================

def get_detected_sections(
    sections: dict[str, list[str]],
) -> list[str]:

    return list(sections.keys())


# ============================================================
# Main Resume Parser
# ============================================================

def parse_resume_text(
    text: str,
) -> dict:

    # --------------------------------------------------------
    # 1. Normalize
    # --------------------------------------------------------

    text = normalize_text(text)

    # --------------------------------------------------------
    # 2. Extract sections
    # --------------------------------------------------------

    sections = extract_sections(
        text
    )

    # --------------------------------------------------------
    # 3. Basic contact
    # --------------------------------------------------------

    full_name = extract_name(
        text
    )

    email = extract_email(
        text
    )

    phone = extract_phone(
        text
    )

    # --------------------------------------------------------
    # 4. Skills
    # --------------------------------------------------------

    skills = extract_skills(
        text
    )

    categorized_skills = (
        extract_categorized_skills(
            text
        )
    )

    # --------------------------------------------------------
    # 5. Structured sections
    # --------------------------------------------------------

    education = extract_education(
        text=text,
        sections=sections,
    )

    experience = extract_experience(
        sections
    )

    projects = extract_projects(
        sections
    )

    certifications = extract_certifications(
        sections
    )

    achievements = extract_achievements(
        sections
    )

    positions = extract_positions(
        sections
    )

    extracurricular = (
        extract_extracurricular(
            sections
        )
    )

    coursework = extract_coursework(
        sections
    )

    # --------------------------------------------------------
    # 6. Academic profile
    # --------------------------------------------------------

    academic_profile = (
        extract_academic_profile(
            text
        )
    )

    # --------------------------------------------------------
    # 7. Explicit experience years
    # --------------------------------------------------------

    total_experience_years = (
        extract_experience_years(
            text
        )
    )

    # --------------------------------------------------------
    # 8. NLP entities
    # --------------------------------------------------------

    named_entities = (
        extract_named_entities(
            text
        )
    )

    organizations = (
        extract_organizations(
            text
        )
    )

    people = (
        extract_people(
            text
        )
    )

    # --------------------------------------------------------
    # 9. Build final parsed object
    # --------------------------------------------------------

    parsed_data = {
        "contact": {
            "full_name": full_name,
            "email": email,
            "phone": phone,
        },

        "summary": extract_summary(
            text
        ),

        "academic_profile": academic_profile,

        "total_experience_years":
            total_experience_years,

        "skills": skills,

        "categorized_skills":
            categorized_skills,

        "education": education,

        "experience": experience,

        "projects": projects,

        "certifications":
            certifications,

        "achievements":
            achievements,

        "positions_of_responsibility":
            positions,

        "extracurricular_activities":
            extracurricular,

        "relevant_coursework":
            coursework,

        "organizations":
            organizations,

        "people":
            people,

        "named_entities":
            named_entities,

        "detected_sections":
            get_detected_sections(
                sections
            ),
    }

    return parsed_data


# ============================================================
# Database: Get Parsed Resume
# ============================================================

def get_parsed_resume(
    db: Session,
    resume_id: uuid.UUID,
) -> ParsedResume | None:

    return (
        db.query(ParsedResume)
        .filter(
            ParsedResume.resume_id
            == resume_id
        )
        .first()
    )


# ============================================================
# Database: Save Parsed Resume
# ============================================================

def save_parsed_resume(
    db: Session,
    resume: Resume,
    parsed_data: dict,
    raw_text: str,
) -> ParsedResume:

    parsed_resume = get_parsed_resume(
        db=db,
        resume_id=resume.id,
    )

    if parsed_resume is None:

        parsed_resume = ParsedResume(
            resume_id=resume.id,
        )

        db.add(parsed_resume)

    contact_data = parsed_data.get(
        "contact",
        {},
    )

    parsed_resume.full_name = (
        contact_data.get(
            "full_name"
        )
    )

    parsed_resume.email = (
        contact_data.get(
            "email"
        )
    )

    parsed_resume.phone = (
        contact_data.get(
            "phone"
        )
    )

    parsed_resume.summary = (
        parsed_data.get(
            "summary"
        )
    )

    parsed_resume.total_experience_years = (
        parsed_data.get(
            "total_experience_years"
        )
    )

    parsed_resume.raw_text = raw_text

    parsed_resume.parsed_data = (
        parsed_data
    )

    db.commit()

    db.refresh(
        parsed_resume
    )

    return parsed_resume


# ============================================================
# Main Resume Processing Pipeline
# ============================================================

def process_resume(
    db: Session,
    resume: Resume,
) -> ParsedResume:

    # --------------------------------------------------------
    # Mark resume as PROCESSING
    # --------------------------------------------------------

    resume.processing_status = (
        ResumeProcessingStatus.PROCESSING
    )

    db.commit()

    db.refresh(resume)

    try:

        # ----------------------------------------------------
        # Locate physical file
        # ----------------------------------------------------

        file_path = Path(
            resume.file_url
        )

        if not file_path.exists():

            raise FileNotFoundError(
                f"Resume file not found: "
                f"{file_path}"
            )

        # ----------------------------------------------------
        # Extract raw text
        # ----------------------------------------------------

        raw_text = extract_resume_text(
            file_path=file_path,
            file_type=resume.file_type,
        )

        if not raw_text:

            raise ValueError(
                "No text could be extracted "
                "from the resume"
            )

        # ----------------------------------------------------
        # Parse resume
        # ----------------------------------------------------

        parsed_data = parse_resume_text(
            raw_text
        )

        # ----------------------------------------------------
        # Save parsed data
        # ----------------------------------------------------

        parsed_resume = save_parsed_resume(
            db=db,
            resume=resume,
            parsed_data=parsed_data,
            raw_text=raw_text,
        )

        # ----------------------------------------------------
        # Mark processing as completed
        # ----------------------------------------------------

        resume.processing_status = (
            ResumeProcessingStatus.COMPLETED
        )

        resume.processed_at = (
            datetime.now(timezone.utc)
        )

        db.commit()

        db.refresh(resume)

        return parsed_resume

    except Exception:

        # ----------------------------------------------------
        # Mark processing as failed
        # ----------------------------------------------------

        resume.processing_status = (
            ResumeProcessingStatus.FAILED
        )

        db.commit()

        raise